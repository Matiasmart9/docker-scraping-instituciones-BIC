import io
from datetime import datetime
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.institucion import Institucion, EstadoActual, HistorialCambios, SnapshotDiario

MESES_NOMBRE = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

# Cuando se migran o siembran datos en bloque (por ejemplo, la carga inicial del
# sistema) pueden registrarse decenas de altas o bajas el mismo día, que no
# representan movimientos reales sino ruido de inicialización. Un día real de
# operación nunca da de alta o de baja tantas instituciones a la vez, así que
# se descarta cualquier día que supere este umbral para no inflar el reporte.
UMBRAL_DIA_MASIVO = 20


def _filtrar_dias_masivos(items: list, obtener_fecha) -> list:
    conteo_por_dia = defaultdict(int)
    for item in items:
        conteo_por_dia[obtener_fecha(item).strftime("%Y-%m-%d")] += 1
    return [item for item in items if conteo_por_dia[obtener_fecha(item).strftime("%Y-%m-%d")] <= UMBRAL_DIA_MASIVO]


def _agrupar_por_mes_y_anio(items: list, obtener_fecha, anio: str, construir_fila) -> dict:
    anio_actual = str(datetime.now().year)
    anios_set = {obtener_fecha(item).strftime("%Y") for item in items}
    anios_set.add(anio_actual)
    anios_disponibles = sorted(anios_set)

    por_mes = defaultdict(list)
    for item in items:
        if obtener_fecha(item).strftime("%Y") != anio:
            continue
        mes = obtener_fecha(item).strftime("%Y-%m")
        por_mes[mes].append(construir_fila(item))

    meses = []
    detalle = []
    for m in range(1, 13):
        clave = f"{anio}-{m:02d}"
        filas = sorted(por_mes.get(clave, []), key=lambda x: x["fecha"])
        meses.append({
            "mes": clave,
            "nombre_mes": MESES_NOMBRE[m - 1],
            "cantidad": len(filas),
        })
        detalle.extend(filas)

    return {
        "anio": anio,
        "anios_disponibles": anios_disponibles,
        "meses": meses,
        "detalle": detalle,
        "total_anio": len(detalle),
    }


def obtener_reporte_bajas(db: Session, anio: str) -> dict:
    """
    Detecta instituciones que pasaron a estado 'Desvinculada' (desde cualquier
    estado anterior), usando el historial de cambios de estado que el scraper ya
    registra en cada corrida (HistorialCambios) apenas detecta la transición — no
    hace falta ningún backfill ni reconstrucción a partir de backups.

    Agrupa por mes calendario del año pedido (los 12 meses, aunque no haya datos
    en alguno) e incluye, por institución, su 'Última Carga XML' vigente (o
    'Sin Datos' si nunca llegó a registrar una carga real).
    """
    todos_los_cambios = db.query(HistorialCambios).filter(
        HistorialCambios.estado_nuevo.ilike("%DESVINCULAD%")
    ).order_by(HistorialCambios.fecha_deteccion.asc()).all()

    cambios = _filtrar_dias_masivos(todos_los_cambios, lambda c: c.fecha_deteccion)

    institucion_ids = {c.institucion_id for c in cambios}
    instituciones_map = {}
    estados_map = {}
    if institucion_ids:
        for inst in db.query(Institucion).filter(Institucion.id.in_(institucion_ids)).all():
            instituciones_map[inst.id] = inst.nombre
        for e in db.query(EstadoActual).filter(EstadoActual.institucion_id.in_(institucion_ids)).all():
            estados_map[e.institucion_id] = e.fecha_ultima_carga

    def construir_fila(c):
        return {
            "institucion_id": c.institucion_id,
            "nombre": instituciones_map.get(c.institucion_id, "Institución desconocida"),
            "fecha": c.fecha_deteccion.strftime("%Y-%m-%d"),
            "estado_anterior": c.estado_anterior or "-",
            "ultima_carga_xml": estados_map.get(c.institucion_id) or None,
        }

    return _agrupar_por_mes_y_anio(cambios, lambda c: c.fecha_deteccion, anio, construir_fila)


def obtener_reporte_altas(db: Session, anio: str) -> dict:
    """
    Detecta instituciones nuevas (dadas de alta por primera vez en el sistema),
    usando la fecha de creación del catálogo de instituciones (Institucion.creado_el),
    que se fija apenas el scraper ve por primera vez un nombre que no coincide con
    ninguna institución ni alias existente — no hace falta ningún backfill.

    Agrupa por mes calendario del año pedido (los 12 meses, aunque no haya datos
    en alguno) e incluye el estado actual de cada institución.
    """
    todas = db.query(Institucion).order_by(Institucion.creado_el.asc()).all()

    # Cuando BICSA renombra una institución, el scraper ve el nombre nuevo
    # como una institución distinta y le crea una fila propia (con su propia
    # creado_el) hasta que un administrador la unifica con la institución
    # original vía "Resolución de Nombres" (endpoint /unificar), lo que
    # traspasa TODO el historial y los snapshots de la institución vieja a la
    # nueva. Para no perder de vista cuándo ingresó realmente, se usa como
    # "fecha de alta efectiva" la más antigua entre su propia creado_el y
    # cualquier snapshot/cambio de estado heredado de una institución
    # unificada en ella — así una institución sigue apareciendo en el mes en
    # que realmente ingresó, aunque BICSA le haya corregido el nombre después.
    #
    # Esto también resuelve solo el caso de instituciones VIEJAS unificadas
    # por error de nombre (ej. una que ya existía desde la siembra inicial
    # del sistema): su fecha de alta efectiva termina cayendo en ese día
    # masivo de inicialización, que ya se descarta como ruido más abajo — sin
    # necesidad de un caso especial aparte.
    ids = [i.id for i in todas]
    fecha_mas_antigua = {}
    if ids:
        for inst_id, fecha_min in (
            db.query(SnapshotDiario.institucion_id, func.min(SnapshotDiario.fecha_snapshot))
            .filter(SnapshotDiario.institucion_id.in_(ids)).group_by(SnapshotDiario.institucion_id).all()
        ):
            fecha_mas_antigua[inst_id] = fecha_min
        for inst_id, fecha_min in (
            db.query(HistorialCambios.institucion_id, func.min(HistorialCambios.fecha_deteccion))
            .filter(HistorialCambios.institucion_id.in_(ids)).group_by(HistorialCambios.institucion_id).all()
        ):
            if inst_id not in fecha_mas_antigua or fecha_min < fecha_mas_antigua[inst_id]:
                fecha_mas_antigua[inst_id] = fecha_min

    def fecha_alta_efectiva(i):
        heredada = fecha_mas_antigua.get(i.id)
        return min(heredada, i.creado_el) if heredada else i.creado_el

    altas = _filtrar_dias_masivos(todas, fecha_alta_efectiva)

    institucion_ids = [i.id for i in altas]
    estados_map = {}
    if institucion_ids:
        for e in db.query(EstadoActual).filter(EstadoActual.institucion_id.in_(institucion_ids)).all():
            estados_map[e.institucion_id] = e.estado

    def construir_fila(i):
        return {
            "institucion_id": i.id,
            "nombre": i.nombre,
            "fecha": fecha_alta_efectiva(i).strftime("%Y-%m-%d"),
            "estado_actual": estados_map.get(i.id, "Desconocido"),
        }

    return _agrupar_por_mes_y_anio(altas, fecha_alta_efectiva, anio, construir_fila)


def _crear_libro_base(titulo: str, subtitulo: str, headers: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.views.sheetView[0].showGridLines = True

    NAVY_HEADER = "1E293B"
    WHITE_TEXT = "FFFFFF"
    BORDER_COLOR = "CBD5E1"

    estilos = {
        "font_title": Font(name="Calibri", size=16, bold=True, color="0F172A"),
        "font_subtitle": Font(name="Calibri", size=10, italic=True, color="64748B"),
        "font_header": Font(name="Calibri", size=11, bold=True, color=WHITE_TEXT),
        "font_cell": Font(name="Calibri", size=10, color="0F172A"),
        "font_cell_sin_datos": Font(name="Calibri", size=10, italic=True, color="94A3B8"),
        "font_bold": Font(name="Calibri", size=10, bold=True, color="0F172A"),
        "thin_border": Border(
            left=Side(style="thin", color=BORDER_COLOR),
            right=Side(style="thin", color=BORDER_COLOR),
            top=Side(style="thin", color=BORDER_COLOR),
            bottom=Side(style="thin", color=BORDER_COLOR),
        ),
    }

    ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    ws["A1"] = titulo
    ws["A1"].font = estilos["font_title"]

    ws.merge_cells(f"A2:{get_column_letter(len(headers))}2")
    ws["A2"] = subtitulo
    ws["A2"].font = estilos["font_subtitle"]

    row_idx = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.font = estilos["font_header"]
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = estilos["thin_border"]
    ws.row_dimensions[row_idx].height = 24

    return wb, ws, estilos


def _finalizar_libro(wb, ws, mensaje_vacio: str | None):
    if mensaje_vacio:
        ws.cell(row=5, column=1, value=mensaje_vacio).font = Font(name="Calibri", size=10, italic=True, color="94A3B8")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
    ws.column_dimensions["A"].width = 42

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generar_excel_bajas(detalle: list[dict], anio: str) -> io.BytesIO:
    """Genera un Excel estilizado con el detalle de instituciones desvinculadas en el año pedido."""
    headers = ["Institución", "Mes", "Fecha de Desvinculación", "Estado Anterior", "Última Carga XML"]
    wb, ws, e = _crear_libro_base(
        f"Desvinculadas {anio}",
        f"Reporte generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M:%S')}. "
        f"Instituciones que pasaron a estado 'Desvinculada' durante el año {anio}.",
        headers,
    )

    row_idx = 5
    for fila in detalle:
        mes_idx = int(fila["fecha"][5:7])
        mes_nombre = f"{MESES_NOMBRE[mes_idx - 1]} {anio}"
        fecha_fmt = datetime.strptime(fila["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y")

        ws.cell(row=row_idx, column=1, value=fila["nombre"]).font = e["font_bold"]
        ws.cell(row=row_idx, column=2, value=mes_nombre).font = e["font_cell"]
        ws.cell(row=row_idx, column=3, value=fecha_fmt).font = e["font_cell"]
        ws.cell(row=row_idx, column=4, value=fila["estado_anterior"]).font = e["font_cell"]

        c_ultima = ws.cell(row=row_idx, column=5)
        if fila["ultima_carga_xml"]:
            c_ultima.value = fila["ultima_carga_xml"]
            c_ultima.font = e["font_cell"]
        else:
            c_ultima.value = "Sin Datos"
            c_ultima.font = e["font_cell_sin_datos"]

        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
            ws.cell(row=row_idx, column=col_idx).border = e["thin_border"]
        row_idx += 1

    return _finalizar_libro(wb, ws, "Sin instituciones desvinculadas en el año seleccionado." if not detalle else None)


def generar_excel_altas(detalle: list[dict], anio: str) -> io.BytesIO:
    """Genera un Excel estilizado con el detalle de instituciones nuevas (altas) en el año pedido."""
    headers = ["Institución", "Mes", "Fecha de Alta", "Estado Actual"]
    wb, ws, e = _crear_libro_base(
        f"Altas {anio}",
        f"Reporte generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M:%S')}. "
        f"Instituciones nuevas registradas durante el año {anio}.",
        headers,
    )

    row_idx = 5
    for fila in detalle:
        mes_idx = int(fila["fecha"][5:7])
        mes_nombre = f"{MESES_NOMBRE[mes_idx - 1]} {anio}"
        fecha_fmt = datetime.strptime(fila["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y")

        ws.cell(row=row_idx, column=1, value=fila["nombre"]).font = e["font_bold"]
        ws.cell(row=row_idx, column=2, value=mes_nombre).font = e["font_cell"]
        ws.cell(row=row_idx, column=3, value=fecha_fmt).font = e["font_cell"]
        ws.cell(row=row_idx, column=4, value=fila["estado_actual"]).font = e["font_cell"]

        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
            ws.cell(row=row_idx, column=col_idx).border = e["thin_border"]
        row_idx += 1

    return _finalizar_libro(wb, ws, "Sin instituciones nuevas en el año seleccionado." if not detalle else None)
